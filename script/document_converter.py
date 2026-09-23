"""
Konverter Dokumen - fungsi-fungsi untuk mengonversi dokumen antar format.

Format didukung: TXT, PDF, DOCX, HTML, MD

Instalasi dependensi (jalankan di terminal / PowerShell):
    pip install -r requirements.txt

DOCX -> PDF berkualitas terbaik butuh Microsoft Word terpasang (lewat docx2pdf).
PDF -> DOCX berkualitas terbaik butuh pustaka pdf2docx.
Jika keduanya tidak tersedia, dipakai jalur cadangan berbasis teks biasa.
"""

import html
import os
import re

from common import unique_output_path

try:
    import pypdf
    PDF_READ_AVAILABLE = True
except ImportError:
    PDF_READ_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    PDF_WRITE_AVAILABLE = True
except ImportError:
    PDF_WRITE_AVAILABLE = False

try:
    import markdown as md_lib
    MARKDOWN_AVAILABLE = True
except ImportError:
    MARKDOWN_AVAILABLE = False

try:
    import html2text
    HTML2TEXT_AVAILABLE = True
except ImportError:
    HTML2TEXT_AVAILABLE = False

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False

try:
    from pdf2docx import Converter as PDF2DocxConverter
    PDF2DOCX_AVAILABLE = True
except ImportError:
    PDF2DOCX_AVAILABLE = False

try:
    from docx2pdf import convert as docx2pdf_convert
    DOCX2PDF_AVAILABLE = True
except ImportError:
    DOCX2PDF_AVAILABLE = False


SUPPORTED_OUTPUT_FORMATS = ["TXT", "PDF", "DOCX", "HTML", "MD"]
SUPPORTED_INPUT_EXTENSIONS = {".txt", ".pdf", ".docx", ".html", ".htm", ".md", ".markdown"}

EXT_TO_FORMAT = {
    ".txt": "TXT",
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".html": "HTML",
    ".htm": "HTML",
    ".md": "MD",
    ".markdown": "MD",
}

FORMAT_TO_EXT = {
    "TXT": ".txt",
    "PDF": ".pdf",
    "DOCX": ".docx",
    "HTML": ".html",
    "MD": ".md",
}


def detect_format(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    return EXT_TO_FORMAT.get(ext, ext.lstrip(".").upper() or "?")


# ---------- Readers: mengubah dokumen sumber menjadi teks biasa ----------

def _read_txt(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_pdf(path):
    if not PDF_READ_AVAILABLE:
        raise RuntimeError("Untuk membaca PDF, install dulu: pip install pypdf")
    reader = pypdf.PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path):
    if not DOCX_AVAILABLE:
        raise RuntimeError("Untuk membaca DOCX, install dulu: pip install python-docx")
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _read_html(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    if HTML2TEXT_AVAILABLE:
        return html2text.html2text(content)
    if BS4_AVAILABLE:
        return BeautifulSoup(content, "html.parser").get_text()
    return re.sub(r"<[^>]+>", "", content)  # fallback kasar: buang tag html


def _load_as_text(path, source_format):
    if source_format == "TXT":
        return _read_txt(path)
    if source_format == "PDF":
        return _read_pdf(path)
    if source_format == "DOCX":
        return _read_docx(path)
    if source_format == "HTML":
        return _read_html(path)
    if source_format == "MD":
        return _read_txt(path)
    raise RuntimeError(f"Format sumber tidak didukung: {source_format}")


# ---------- Writers: menulis teks biasa menjadi format tujuan ----------

def _write_txt(text, output_path):
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)


def _write_docx(text, output_path):
    if not DOCX_AVAILABLE:
        raise RuntimeError("Untuk menulis DOCX, install dulu: pip install python-docx")
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(output_path)


def _write_pdf(text, output_path):
    if not PDF_WRITE_AVAILABLE:
        raise RuntimeError("Untuk menulis PDF, install dulu: pip install reportlab")
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(output_path, pagesize=A4)
    story = [
        Paragraph(html.escape(line) if line.strip() else "&nbsp;", styles["Normal"])
        for line in text.split("\n")
    ]
    doc.build(story)


def _write_html(text, output_path, source_format):
    if source_format == "MD" and MARKDOWN_AVAILABLE:
        body = md_lib.markdown(text)
    else:
        escaped = html.escape(text)
        body = "".join(f"<p>{line}</p>\n" for line in escaped.split("\n"))
    content = (
        '<!DOCTYPE html>\n<html>\n<head><meta charset="utf-8"></head>\n'
        f"<body>\n{body}\n</body>\n</html>\n"
    )
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)


def _write_md(text, output_path, source_format):
    if source_format == "HTML" and HTML2TEXT_AVAILABLE:
        text = html2text.html2text(text)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)


def convert_document(input_path, output_dir, target_format):
    """Konversi satu file dokumen. Mengembalikan path file hasil."""
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    source_format = detect_format(input_path)
    target_format = target_format.upper()
    ext = FORMAT_TO_EXT.get(target_format, f".{target_format.lower()}")
    output_path = unique_output_path(output_dir, base_name, ext)

    # Jalur khusus dengan kualitas lebih baik bila pustaka tersedia
    if source_format == "PDF" and target_format == "DOCX" and PDF2DOCX_AVAILABLE:
        converter = PDF2DocxConverter(input_path)
        converter.convert(output_path)
        converter.close()
        return output_path

    if source_format == "DOCX" and target_format == "PDF" and DOCX2PDF_AVAILABLE:
        try:
            docx2pdf_convert(input_path, output_path)
            return output_path
        except Exception:
            pass  # fallback ke jalur teks biasa di bawah

    if source_format == target_format:
        with open(input_path, "rb") as src, open(output_path, "wb") as dst:
            dst.write(src.read())
        return output_path

    text = _load_as_text(input_path, source_format)

    if target_format == "TXT":
        _write_txt(text, output_path)
    elif target_format == "DOCX":
        _write_docx(text, output_path)
    elif target_format == "PDF":
        _write_pdf(text, output_path)
    elif target_format == "HTML":
        _write_html(text, output_path, source_format)
    elif target_format == "MD":
        _write_md(text, output_path, source_format)
    else:
        raise RuntimeError(f"Format tujuan tidak didukung: {target_format}")

    return output_path
